from __future__ import annotations

from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from psycopg import AsyncConnection

from app.core.audit import record_audit
from app.core.cache import PRODUCT_PREFIX, get_json, invalidate_prefix, set_json
from app.core.errors import ConflictError, NotFoundError, ValidationAppError
from app.core.text import slugify
from app.repositories import category_repo, inventory_repo, product_repo, variant_repo
from app.schemas.common import Page
from app.schemas.product import ExcelImportResult, ProductCreate, ProductUpdate, VariantCreate, VariantUpdate

ADMIN_STAFF_ROLES = ("admin", "staff")


def _is_privileged(current_user: dict[str, Any] | None) -> bool:
    return current_user is not None and current_user["role"] in ADMIN_STAFF_ROLES


async def _add_variant(conn: AsyncConnection, product_id: int, variant: VariantCreate) -> None:
    if await variant_repo.sku_exists(conn, variant.sku):
        raise ConflictError(f"SKU biến thể '{variant.sku}' đã tồn tại")
    row = await variant_repo.create(
        conn,
        product_id=product_id,
        variant_name=variant.variant_name,
        sku=variant.sku,
        price_adjustment=variant.price_adjustment,
        is_active=variant.is_active,
    )
    await inventory_repo.create_for_variant(
        conn,
        product_variant_id=row["id"],
        quantity_on_hand=variant.initial_quantity,
        reorder_level=variant.reorder_level,
    )


async def create_product(conn: AsyncConnection, data: ProductCreate, current_user: dict[str, Any]) -> dict[str, Any]:
    if await category_repo.get_by_id(conn, data.category_id) is None:
        raise NotFoundError(f"Không tìm thấy category id={data.category_id}")
    if await product_repo.sku_exists(conn, data.sku):
        raise ConflictError(f"SKU '{data.sku}' đã tồn tại")

    slug = slugify(data.slug or data.name)
    if await product_repo.slug_exists(conn, slug):
        slug = f"{slug}-{data.sku.lower()}"

    product_id = await product_repo.create(
        conn,
        category_id=data.category_id,
        sku=data.sku,
        name=data.name,
        slug=slug,
        author=data.author,
        publisher=data.publisher,
        isbn=data.isbn,
        description=data.description,
        cover_image_url=data.cover_image_url,
        base_price=data.base_price,
        is_active=data.is_active,
    )

    for variant in data.variants:
        await _add_variant(conn, product_id, variant)

    product = await product_repo.get_by_id(conn, product_id)
    assert product is not None
    await record_audit(
        conn, user_id=current_user["id"], action="create", entity_type="product",
        entity_id=product_id, new_value=product,
    )
    await invalidate_prefix(PRODUCT_PREFIX)
    return product


async def list_products(
    conn: AsyncConnection,
    *,
    keyword: str | None,
    category_id: int | None,
    min_price: float | None,
    max_price: float | None,
    is_active: bool | None,
    sort_by: str,
    sort_dir: str,
    limit: int,
    offset: int,
    current_user: dict[str, Any] | None,
) -> Page:
    # Khách vãng lai/customer chỉ thấy sản phẩm đang bán, bất kể query param truyền gì
    effective_is_active = is_active if _is_privileged(current_user) else True

    cache_key = (
        f"{PRODUCT_PREFIX}:list:{keyword}:{category_id}:{min_price}:{max_price}:"
        f"{effective_is_active}:{sort_by}:{sort_dir}:{limit}:{offset}"
    )
    cached = await get_json(cache_key)
    if cached is not None:
        return Page(**cached)

    rows, total = await product_repo.list_products(
        conn,
        keyword=keyword,
        category_id=category_id,
        min_price=min_price,
        max_price=max_price,
        is_active=effective_is_active,
        sort_by=sort_by,
        sort_dir=sort_dir,
        limit=limit,
        offset=offset,
    )
    items = [{k: v for k, v in row.items() if k != "total_count"} for row in rows]
    page = Page(items=items, total=total, limit=limit, offset=offset)
    await set_json(cache_key, page.model_dump())
    return page


async def get_product(
    conn: AsyncConnection, product_id: int, current_user: dict[str, Any] | None
) -> dict[str, Any]:
    product = await product_repo.get_by_id(conn, product_id)
    if product is None:
        raise NotFoundError(f"Không tìm thấy sản phẩm id={product_id}")
    if not product["is_active"] and not _is_privileged(current_user):
        raise NotFoundError(f"Không tìm thấy sản phẩm id={product_id}")
    return product


async def update_product(
    conn: AsyncConnection, product_id: int, data: ProductUpdate, current_user: dict[str, Any]
) -> dict[str, Any]:
    """update_product() (DB) ghi đè toàn bộ giá trị — merge field cũ + field
    client gửi (exclude_unset) thành giá trị cuối cùng trước khi gọi."""
    old = await get_product(conn, product_id, current_user)
    payload = data.model_dump(exclude_unset=True)

    new_category_id = payload.get("category_id", old["category_id"])
    if "category_id" in payload and await category_repo.get_by_id(conn, new_category_id) is None:
        raise NotFoundError(f"Không tìm thấy category id={new_category_id}")

    if "slug" in payload:
        new_slug = slugify(payload["slug"])
    elif "name" in payload:
        new_slug = slugify(payload["name"])
    else:
        new_slug = old["slug"]

    if new_slug != old["slug"] and await product_repo.slug_exists(conn, new_slug):
        raise ConflictError(f"Slug '{new_slug}' đã tồn tại")

    merged = {
        "category_id": new_category_id,
        "slug": new_slug,
        "name": payload.get("name", old["name"]),
        "author": payload.get("author", old["author"]),
        "publisher": payload.get("publisher", old["publisher"]),
        "isbn": payload.get("isbn", old["isbn"]),
        "description": payload.get("description", old["description"]),
        "cover_image_url": payload.get("cover_image_url", old["cover_image_url"]),
        "base_price": payload.get("base_price", old["base_price"]),
        "is_active": payload.get("is_active", old["is_active"]),
    }

    updated_ok = await product_repo.update(conn, product_id, **merged)
    if not updated_ok:
        raise NotFoundError(f"Không tìm thấy sản phẩm id={product_id}")

    updated = await product_repo.get_by_id(conn, product_id)
    assert updated is not None
    await record_audit(
        conn, user_id=current_user["id"], action="update", entity_type="product",
        entity_id=product_id, old_value=old, new_value=updated,
    )
    await invalidate_prefix(PRODUCT_PREFIX)
    return updated


async def delete_product(conn: AsyncConnection, product_id: int, current_user: dict[str, Any]) -> None:
    old = await get_product(conn, product_id, current_user)
    deleted = await product_repo.soft_delete(conn, product_id)
    if not deleted:
        raise NotFoundError(f"Không tìm thấy sản phẩm id={product_id}")
    await record_audit(
        conn, user_id=current_user["id"], action="delete", entity_type="product",
        entity_id=product_id, old_value=old,
    )
    await invalidate_prefix(PRODUCT_PREFIX)


async def add_variant(
    conn: AsyncConnection, product_id: int, data: VariantCreate, current_user: dict[str, Any]
) -> dict[str, Any]:
    product = await product_repo.get_by_id(conn, product_id)
    if product is None:
        raise NotFoundError(f"Không tìm thấy sản phẩm id={product_id}")

    await _add_variant(conn, product_id, data)
    await record_audit(
        conn, user_id=current_user["id"], action="create", entity_type="product_variant",
        entity_id=product_id, new_value={"sku": data.sku, "variant_name": data.variant_name},
    )
    await invalidate_prefix(PRODUCT_PREFIX)

    refreshed = await product_repo.get_by_id(conn, product_id)
    assert refreshed is not None
    return refreshed


async def update_variant(
    conn: AsyncConnection, variant_id: int, data: VariantUpdate, current_user: dict[str, Any]
) -> dict[str, Any]:
    """update_variant() (DB) ghi đè toàn bộ giá trị — merge field cũ + field
    client gửi (exclude_unset) thành giá trị cuối cùng trước khi gọi."""
    old = await variant_repo.get_by_id(conn, variant_id)
    if old is None:
        raise NotFoundError(f"Không tìm thấy biến thể id={variant_id}")

    payload = data.model_dump(exclude_unset=True, exclude={"reorder_level"})
    updated = await variant_repo.update(
        conn,
        variant_id,
        variant_name=payload.get("variant_name", old["variant_name"]),
        price_adjustment=payload.get("price_adjustment", old["price_adjustment"]),
        is_active=payload.get("is_active", old["is_active"]),
    )

    if data.reorder_level is not None:
        await inventory_repo.update_reorder_level(conn, variant_id, data.reorder_level)

    await record_audit(
        conn, user_id=current_user["id"], action="update", entity_type="product_variant",
        entity_id=variant_id, old_value=old, new_value=updated,
    )
    await invalidate_prefix(PRODUCT_PREFIX)

    refreshed = await product_repo.get_by_id(conn, old["product_id"])
    assert refreshed is not None
    return refreshed


async def deactivate_variant(conn: AsyncConnection, variant_id: int, current_user: dict[str, Any]) -> dict[str, Any]:
    old = await variant_repo.get_by_id(conn, variant_id)
    if old is None:
        raise NotFoundError(f"Không tìm thấy biến thể id={variant_id}")

    updated = await variant_repo.set_active(conn, variant_id, False)
    await record_audit(
        conn, user_id=current_user["id"], action="deactivate", entity_type="product_variant",
        entity_id=variant_id, old_value=old, new_value=updated,
    )
    await invalidate_prefix(PRODUCT_PREFIX)

    refreshed = await product_repo.get_by_id(conn, old["product_id"])
    assert refreshed is not None
    return refreshed


_EXCEL_REQUIRED_HEADERS = ["sku", "name", "category_id", "base_price"]
_EXCEL_OPTIONAL_HEADERS = ["author", "publisher", "isbn", "description", "cover_image_url", "slug"]


async def import_products_from_excel(
    conn: AsyncConnection, file_bytes: bytes, current_user: dict[str, Any]
) -> ExcelImportResult:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ValidationAppError("Thiếu thư viện openpyxl trên server") from exc

    try:
        workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ValidationAppError(f"File Excel không hợp lệ: {exc}") from exc

    sheet = workbook.active
    rows_iter = sheet.iter_rows(values_only=True)
    try:
        header = [str(h).strip().lower() if h is not None else "" for h in next(rows_iter)]
    except StopIteration:
        raise ValidationAppError("File Excel rỗng") from None

    missing = [h for h in _EXCEL_REQUIRED_HEADERS if h not in header]
    if missing:
        raise ValidationAppError(f"Thiếu cột bắt buộc trong Excel: {', '.join(missing)}")

    col_index = {name: header.index(name) for name in header if name}

    created = 0
    errors: list[str] = []

    for row_number, row in enumerate(rows_iter, start=2):
        if row is None or all(cell is None for cell in row):
            continue

        def cell(name: str, _row: tuple = row) -> Any:
            idx = col_index.get(name)
            return _row[idx] if idx is not None and idx < len(_row) else None

        try:
            # Mỗi dòng chạy trong 1 SAVEPOINT riêng (conn.transaction() lồng trong
            # transaction của request): lỗi ở 1 dòng chỉ rollback dòng đó, không
            # làm hỏng transaction chung và không chặn các dòng còn lại.
            async with conn.transaction():
                sku = str(cell("sku") or "").strip()
                name = str(cell("name") or "").strip()
                category_id_raw = cell("category_id")
                base_price_raw = cell("base_price")

                if not sku or not name or category_id_raw is None or base_price_raw is None:
                    raise ValueError("thiếu sku/name/category_id/base_price")

                category_id = int(category_id_raw)
                try:
                    base_price = Decimal(str(base_price_raw))
                except InvalidOperation as exc:
                    raise ValueError(f"base_price không hợp lệ: {base_price_raw}") from exc

                if await category_repo.get_by_id(conn, category_id) is None:
                    raise ValueError(f"category_id {category_id} không tồn tại")
                if await product_repo.sku_exists(conn, sku):
                    raise ValueError(f"sku '{sku}' đã tồn tại")

                slug_raw = cell("slug")
                slug = slugify(str(slug_raw)) if slug_raw else slugify(name)
                if await product_repo.slug_exists(conn, slug):
                    slug = f"{slug}-{sku.lower()}"

                product_id = await product_repo.create(
                    conn,
                    category_id=category_id,
                    sku=sku,
                    name=name,
                    slug=slug,
                    author=str(cell("author")) if cell("author") else None,
                    publisher=str(cell("publisher")) if cell("publisher") else None,
                    isbn=str(cell("isbn")) if cell("isbn") else None,
                    description=str(cell("description")) if cell("description") else None,
                    cover_image_url=str(cell("cover_image_url")) if cell("cover_image_url") else None,
                    base_price=base_price,
                    is_active=True,
                )
                await record_audit(
                    conn, user_id=current_user["id"], action="create", entity_type="product",
                    entity_id=product_id, new_value={"sku": sku, "name": name, "source": "excel_import"},
                )
            created += 1
        except Exception as exc:  # noqa: BLE001 - gom lỗi từng dòng, không dừng cả file
            errors.append(f"Dòng {row_number}: {exc}")

    if created:
        await invalidate_prefix(PRODUCT_PREFIX)

    return ExcelImportResult(created=created, failed=len(errors), errors=errors)
